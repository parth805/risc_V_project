import pptx

pptx_path = r'C:/Users/Asus/.gemini/antigravity/brain/a927e01c-747b-4752-ac10-87924e02d758/.user_uploaded/media_1792953919650_b6db8f41.pptx'
prs = pptx.Presentation(pptx_path)

for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for s_idx, shape in enumerate(slide.shapes):
        pos_str = f"L={shape.left/914400:.2f}in, T={shape.top/914400:.2f}in, W={shape.width/914400:.2f}in, H={shape.height/914400:.2f}in"
        print(f'  Shape {s_idx}: {shape.name} ({pos_str}) type={shape.shape_type}')
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                p_text = p.text.strip().replace('\n', ' ')
                font_name = p.font.name if p.font else 'None'
                font_size = p.font.size.pt if (p.font and p.font.size) else 'None'
                font_bold = p.font.bold if p.font else 'None'
                if p_text:
                    print(f'    P{p_idx}: "{p_text[:80]}" (font={font_name}, size={font_size}, bold={font_bold})')
