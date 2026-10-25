import pptx

prs = pptx.Presentation(r'C:/Users/Asus/.gemini/antigravity/brain/a927e01c-747b-4752-ac10-87924e02d758/.user_uploaded/media_1792953919650_b6db8f41.pptx')

for i in range(min(3, len(prs.slides))):
    slide = prs.slides[i]
    print(f'=== Slide {i+1} ===')
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    color = r.font.color.rgb if (r.font and r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB) else 'None'
                    print(f'  Text: "{r.text[:30]}" | Font: {r.font.name} | Size: {r.font.size.pt if r.font.size else "None"} | Bold: {r.font.bold} | Color: {color}')
