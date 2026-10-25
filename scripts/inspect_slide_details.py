import pptx

prs = pptx.Presentation(r'C:/Users/Asus/.gemini/antigravity/brain/a927e01c-747b-4752-ac10-87924e02d758/.user_uploaded/media_1792953919650_b6db8f41.pptx')

for i, slide in enumerate(prs.slides):
    print(f'=== Slide {i+1} ===')
    for s in slide.shapes:
        if s.has_text_frame:
            first_line = s.text_frame.text.splitlines()[0] if s.text_frame.text.splitlines() else ''
            print(f'  ID={s.shape_id}, name="{s.name}", pos=({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f}), text="{first_line[:40]}"')
