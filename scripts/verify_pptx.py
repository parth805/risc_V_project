import pptx

prs = pptx.Presentation(r'P:/project 1/RISC_V_AI_Accelerator_Review3_Presentation.pptx')
print('Total slides:', len(prs.slides))

for i, slide in enumerate(prs.slides):
    print(f'=== Slide {i+1} ===')
    for s in slide.shapes:
        s_type = s.shape_type
        pos = f'({s.left/914400:.2f}, {s.top/914400:.2f}, {s.width/914400:.2f}, {s.height/914400:.2f})'
        if s.has_text_frame:
            t = s.text_frame.text.strip().replace('\n', ' ')
            print(f'  [Text] {s.name} at {pos}: "{t[:60]}"')
        else:
            print(f'  [Other] {s.name} type={s_type} at {pos}')
